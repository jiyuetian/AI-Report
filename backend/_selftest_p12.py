"""P1-2 表头探测打分自测：合成"多行说明型表格"，验证能正确探测真实列头行。"""
import sys
from pathlib import Path

sys.path.insert(0, r"C:/Users/Asus009/Desktop/临时/ai大赛/AI-Report/backend")

import openpyxl
from app.core.file_parser import FileParser

BASE = Path(r"C:/Users/Asus009/Desktop/临时/ai大赛/AI-Report/backend")

# ---- 用例1：多行说明型表格（首行合并标题、次行说明、第3行才是列头） ----
p1 = BASE / "_p12_test_multi.xlsx"
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "明细分页"
ws["A1"] = "贷款明细表（2025年度）"              # 行0：合并标题（仅首格）
ws["A2"] = "注：本表统计各笔贷款基本画像信息，单位：元"  # 行1：说明（仅首格，含"注："和"单位："）
hdr = ["贷款编号", "客户编号", "区域", "贷款金额", "状态"]   # 行2：真实列头
for j, h in enumerate(hdr, 1):
    ws.cell(row=3, column=j, value=h)
data = [
    ["DB001", "C001", "西南", 120000, "正常"],
    ["DB002", "C002", "华北", 88000, "逾期"],
    ["DB003", "C003", "华东", 150000, "正常"],
    ["DB004", "C004", "华南", 99000, "结清"],
]
for i, r in enumerate(data, 4):
    for j, v in enumerate(r, 1):
        ws.cell(row=i, column=j, value=v)
wb.save(p1)

idx1 = FileParser._detect_header_row(p1, sheet_index=0)
print(f"[用例1] detected header_row_index = {idx1} (期望 2)")
res1 = FileParser.parse_file(p1, ".xlsx")
print(f"[用例1] parse header_row_index = {res1.get('header_row_index')}")
print(f"[用例1] columns = {list(res1['dataframe'].columns)}")
print(f"[用例1] rows = {len(res1['dataframe'])}")
assert idx1 == 2, "P1-2 用例1 探测失败"
assert list(res1["dataframe"].columns) == hdr, "P1-2 用例1 列名错误"
assert len(res1["dataframe"]) == 4, "P1-2 用例1 行数错误"

# ---- 用例2：普通首行即列头（应回退到行0，行为不变） ----
p2 = BASE / "_p12_test_normal.xlsx"
wb2 = openpyxl.Workbook()
ws2 = wb2.active
for j, h in enumerate(["编号", "姓名", "金额"], 1):
    ws2.cell(row=1, column=j, value=h)
for i, r in enumerate([["A1", "张三", 100], ["A2", "李四", 200]], 2):
    for j, v in enumerate(r, 1):
        ws2.cell(row=i, column=j, value=v)
wb2.save(p2)
idx2 = FileParser._detect_header_row(p2, sheet_index=0)
print(f"[用例2] detected header_row_index = {idx2} (期望 0)")
res2 = FileParser.parse_file(p2, ".xlsx")
assert idx2 == 0, "P1-2 用例2 应回退首行"
assert list(res2["dataframe"].columns) == ["编号", "姓名", "金额"], "P1-2 用例2 列名错误"

p1.unlink()
p2.unlink()
print("P1-2 PASS")
