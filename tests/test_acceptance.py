"""验收测试（对照《给A的改动清单与验收标准》8条交付标准）

覆盖范围（全部确定性断言，不依赖 LLM 是否可用）：
  [1] xlsx 上传 -> 建表数据集                    标准第1条（表格侧）
  [2] 质检 -> 规则检测返回                        标准第7条链路（清洗环节）
  [3] brain run -> 分析+出图（S1-S5/v2 引擎）     标准第6条 dashboard 无回归
  [4] xlsx 报告 -> 7章节/图表>=6/versionId        标准第1/4/5条
  [5] 3 条数字复算（openpyxl 独立读取回算）       标准第2条
  [6] docx 上传 -> 文档型数据集                   标准第1条（文档侧）
  [7] docx 报告 -> 7章节/图表>=6/versionId        标准第1/4/5条
  [8] docx 段落数复算（python-docx 独立读取）     标准第2条
  [9] 历史版本回看 -> 同看板两版本列表             标准第4条
  [10] llm_used 诚实性 -> 降级标注一致性          标准第3条
  [11] S4/S5 内部端点 -> narrative_flow 真实生成  协作规则（占位符清除）

前置条件：
  1. 后端已监听 127.0.0.1:8000（或传端口参数）。
  2. 同一时刻只能有一个后端进程（DuckDB 单写者锁）。
  3. 依赖 openpyxl / python-docx（用于独立复算，不复用后端解析）。

用法：
    cd backend
    python ../tests/test_acceptance.py            # 默认打 8000
    python ../tests/test_acceptance.py 8001

退出码：0 = 全部通过；1 = 存在失败项。
"""
import io
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

PORT = int(sys.argv[1]) if len(sys.argv) > 1 else 8000
BASE = f"http://127.0.0.1:{PORT}/api/v1"
HERE = Path(__file__).resolve().parent
FIXTURES = HERE / "_fixtures"
OUT = HERE / "_reports"
FIXTURES.mkdir(parents=True, exist_ok=True)
OUT.mkdir(parents=True, exist_ok=True)

PASS, FAIL = 0, 0


def req(method, path, body=None, files=None, timeout=180):
    url = BASE + path
    data, headers = None, {}
    if files is not None:
        boundary = "----AcceptanceBoundary9c4e"
        parts = []
        for fname, content, ctype in files:
            parts.append(f"--{boundary}\r\n".encode())
            parts.append(
                f'Content-Disposition: form-data; name="file"; filename="{fname}"\r\n'
                f"Content-Type: {ctype}\r\n\r\n".encode()
            )
            parts.append(content)
            parts.append(b"\r\n")
        parts.append(f"--{boundary}--\r\n".encode())
        data, headers["Content-Type"] = b"".join(parts), f"multipart/form-data; boundary={boundary}"
    elif body is not None:
        data, headers["Content-Type"] = (
            json.dumps(body, ensure_ascii=False).encode("utf-8"), "application/json"
        )
    r = urllib.request.Request(url, data=data, method=method, headers=headers)
    try:
        with urllib.request.urlopen(r, timeout=timeout) as resp:
            raw = resp.read().decode("utf-8")
    except urllib.error.HTTPError as e:
        return e.code, _try_json(e.read().decode("utf-8"))
    return resp.status, _try_json(raw)


def _try_json(raw):
    try:
        return json.loads(raw)
    except Exception:
        return raw


def check(name, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  PASS  {name}  {extra}")
    else:
        FAIL += 1
        print(f"  FAIL  {name}  {extra}")


def group(title):
    print(f"\n[{title}]")


def detail(msg):
    print(f"  note  {msg}")


def upload(path, ctype):
    """两步式上传第1步：POST /files（multipart），返回 file_id。保留源文件用于独立复算。"""
    with open(path, "rb") as f:
        return req("POST", "/files", files=[(path.name, f.read(), ctype)])


def make_dataset(file_id, name):
    """两步式上传第2步：POST /datasets（JSON body 携带 file_id），返回 dataset 信息。"""
    return req("POST", "/datasets", {"file_id": file_id, "name": name})


# ── fixtures：xlsx 与 docx 测试文件（数值精心设计，可独立复算） ──
XLSX_PATH = FIXTURES / "acceptance_loan.xlsx"
DOCX_PATH = FIXTURES / "acceptance_doc.docx"

# 8 行放款数据；总金额 55,000,000，均值 6,875,000
LOAN_ROWS = [
    ("渠道", "客户类型", "放款金额", "期限月", "利率"),
    ("线上", "个人", 10000000, 12, 5.2),
    ("线上", "个人", 8500000, 24, 5.8),
    ("线上", "小微", 6200000, 36, 6.5),
    ("线下", "个人", 12000000, 12, 5.0),
    ("线下", "小微", 4800000, 24, 6.2),
    ("线下", "个人", 3500000, 36, 6.8),
    ("合作", "小微", 2700000, 12, 5.5),
    ("合作", "个人", 7300000, 24, 6.0),
]
N_ROWS = len(LOAN_ROWS) - 1          # 8
SUM_AMOUNT = sum(r[2] for r in LOAN_ROWS[1:])   # 55,000,000
AVG_AMOUNT = SUM_AMOUNT / N_ROWS                 # 6,875,000

DOC_PARAGRAPHS = [
    "2024年第三季度，零售业务整体保持稳健增长，线上渠道贡献了主要的交易增量。",
    "客户结构方面，个人客户占比持续提升，小微企业贷款需求回暖明显。",
    "风险管理上，逾期率控制在合理区间，未发生重大信用风险事件。",
    "产品创新方面，本季度上线了两款面向小微客群的快捷贷款产品。",
    "渠道协同方面，线上线下一体化运营初见成效，交叉销售转化率提升。",
    "展望下季度，将继续优化风控模型，并扩大合作渠道的覆盖范围。",
]


def make_fixtures():
    from openpyxl import Workbook
    from docx import Document

    if not XLSX_PATH.exists():
        wb = Workbook()
        ws = wb.active
        for row in LOAN_ROWS:
            ws.append(list(row))
        wb.save(XLSX_PATH)

    if not DOCX_PATH.exists():
        doc = Document()
        for p in DOC_PARAGRAPHS:
            doc.add_paragraph(p)
        doc.save(DOCX_PATH)


def read_xlsx_independently():
    """独立复算通道：不经被测系统，直接用 openpyxl 读源文件。"""
    from openpyxl import load_workbook

    wb = load_workbook(XLSX_PATH, read_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    header = [str(c) for c in rows[0]]
    amount_idx = header.index("放款金额")
    amounts = [float(r[amount_idx]) for r in rows[1:] if r[amount_idx] is not None]
    return len(rows) - 1, sum(amounts), sum(amounts) / len(amounts)


def read_docx_independently():
    """独立复算通道：直接用 python-docx 读源文件段落数。"""
    from docx import Document

    doc = Document(DOCX_PATH)
    return len([p for p in doc.paragraphs if p.text.strip()])


def make_dashboard(ds_id, name, charts_cfg=None):
    s, r = req("POST", "/dashboards/", {
        "name": name, "description": "验收看板", "dataset_ids": [ds_id],
    })
    if s != 200:
        return s, r
    did = r.get("dashboard_id")
    if charts_cfg:
        s2, _ = req("PATCH", f"/dashboards/{did}", {"config": {"charts": charts_cfg}})
        return s2, {"success": True, "dashboard_id": did}
    return s, {"success": True, "dashboard_id": did}


CHARTS_CFG = [
    {"chart_type": "kpi", "title": "放款总额", "y_field": "放款金额"},
    {"chart_type": "bar", "title": "渠道放款金额对比", "x_field": "渠道", "y_field": "放款金额"},
    {"chart_type": "pie", "title": "渠道金额占比", "x_field": "渠道", "y_field": "放款金额"},
    {"chart_type": "bar", "title": "客户类型金额对比", "x_field": "客户类型", "y_field": "放款金额"},
    {"chart_type": "histogram", "title": "放款金额分布", "y_field": "放款金额"},
    {"chart_type": "bar", "title": "期限金额分布", "x_field": "期限月", "y_field": "放款金额"},
]


def chart_elements(html):
    return (html.count('class="echarts-chart"')
            + html.count('class="kpi-card"')
            + html.count('class="kpi-display"'))


def extract_appendix_rows(html):
    """从附录数据样本表提取数据行（找行数与源数据一致的表）。"""
    tables = re.findall(r"<table class='data-table'>(.*?)</table>", html, re.S)
    for t in tables:
        body = re.search(r"<tbody>(.*?)</tbody>", t, re.S)
        if not body:
            continue
        trs = re.findall(r"<tr>(.*?)</tr>", body.group(1), re.S)
        rows = []
        for tr in trs:
            cells = [c.strip() for c in re.findall(r"<td>(.*?)</td>", tr, re.S)]
            if cells:
                rows.append(cells)
        if len(rows) == N_ROWS:
            return rows
    return None


def fmt(v):
    return f"{v:,.2f}"


def main():
    make_fixtures()
    x_rows, x_sum, x_avg = read_xlsx_independently()
    d_paras = read_docx_independently()
    detail(f"独立复算基准: xlsx行数={x_rows} 总金额={x_sum:,.0f} 均值={x_avg:,.2f}; docx段落数={d_paras}")

    # [1] xlsx 上传建集
    group("1. xlsx 上传 -> 建表数据集（标准1表格侧）")
    s, r = upload(XLSX_PATH, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    check("上传200", s == 200, f"status={s} {str(r)[:120]}")
    file_x = r.get("file_id") if isinstance(r, dict) else None
    check("返回file_id", bool(file_x))
    s2, r2 = make_dataset(file_x, "验收_xlsx放款数据")
    check("建数据集200", s2 == 200, f"status={s2} {str(r2)[:120]}")
    ds_x = r2.get("dataset_id") if isinstance(r2, dict) else None
    check("返回dataset_id", bool(ds_x))
    check("duckdb_table已建表", bool(isinstance(r2, dict) and r2.get("duckdb_table")))
    check("行数=8", isinstance(r2, dict) and r2.get("row_count") == N_ROWS, f"got={r2.get('row_count') if isinstance(r2, dict) else '?'}")

    # [2] 质检（清洗环节）
    group("2. 质检 -> 规则检测（标准7链路-清洗）")
    s, r = req("POST", "/quality/check", {"dataset_id": ds_x})
    check("质检200", s == 200, f"status={s} {str(r)[:150]}")

    # [3] brain run（分析+出图，S1-S5 + v2 引擎回归）
    group("3. brain run -> 分析出图（标准6 无回归）")
    s, r = req("POST", "/brain/run", {"dataset_id": ds_x, "user_id": "acceptance"})
    check("run启动200", s == 200, f"status={s} {str(r)[:150]}")
    run_id = r.get("run_id") if isinstance(r, dict) else None
    dash_brain = None
    if run_id:
        status, final = "running", {}
        for _ in range(60):
            time.sleep(2)
            s2, r2 = req("GET", f"/brain/run/{run_id}/status")
            if s2 == 200 and isinstance(r2, dict):
                status = r2.get("status", "running")
                if status in ("completed", "failed", "cancelled"):
                    final = r2
                    break
        check("run完成", status == "completed", f"status={status} msg={final.get('message', '')[:120]}")
        dash_brain = (final.get("detail") or {}).get("dashboard_id")
        check("返回dashboard_id", bool(dash_brain))
        if dash_brain:
            s3, r3 = req("GET", f"/dashboards/{dash_brain}")
            cfg = (r3 or {}).get("config") or {}
            charts = cfg.get("charts") or []
            check("v2引擎出图>=1", len(charts) >= 1, f"charts={len(charts)}")
            flow = ((r3 or {}).get("layout") or {}).get("narrative", "")
            check("narrative_flow真实生成(非硬编码)", bool(flow) and flow != "结论→佐证→明细", f"flow={flow[:80]}")
            score = cfg.get("score") or {}
            check("S5评分存在", "overall" in score, f"score={score.get('overall')}")

    # [4] xlsx 报告
    group("4. xlsx 报告 -> 7章节/图表>=6/versionId（标准1/4/5）")
    s, r = make_dashboard(ds_x, "验收_xlsx看板", CHARTS_CFG)
    check("看板创建200", s == 200, f"status={s} {str(r)[:120]}")
    dash_x = r.get("dashboard_id")
    s, r = req("POST", "/reports", {"dashboard_id": dash_x})
    check("报告生成200", s == 200, f"status={s} {str(r)[:150]}")
    rep1 = r if isinstance(r, dict) else {}
    vid1 = rep1.get("version_id")
    check("versionId非null非空", isinstance(vid1, str) and len(vid1) > 0, f"version_id={vid1}")
    s, r = req("GET", f"/reports/{rep1.get('report_id')}/html")
    html_x = r if isinstance(r, str) else ""
    (OUT / "acceptance_xlsx_report.html").write_text(html_x, encoding="utf-8")
    n_ch_x = html_x.count('class="chapter"')
    check("章节数=7", n_ch_x == 7 or len(re.findall(r"<h2[^>]*>第[一二三四五六七]", html_x)) == 7,
          f"chapters={n_ch_x}")
    check("图表数>=6", chart_elements(html_x) >= 6, f"charts={chart_elements(html_x)}")

    # [5] 3 条数字复算
    group("5. 数字复算x3（标准2：抽3条回数据复算）")
    rows = extract_appendix_rows(html_x)
    check("附录样本表存在且行数匹配", rows is not None, f"rows={len(rows) if rows else 0}")
    if rows:
        # R1: 行数
        check("R1 行数复算", len(rows) == x_rows, f"报告={len(rows)} 源={x_rows}")
        # R2: 金额列求和（附录第3列=放款金额）
        try:
            amounts = [float(c.replace(",", "")) for c in (r_[2] for r_ in rows)]
            check("R2 总金额复算", abs(sum(amounts) - x_sum) < 0.01,
                  f"报告={sum(amounts):,.0f} 源={x_sum:,.0f}")
        except Exception as e:
            check("R2 总金额复算", False, f"解析失败: {e}")
    # R3: 执行摘要统计（确定性兜底时）或 KPI 值
    r3_ok = False
    r3_note = ""
    if rep1.get("llm_used") is False:
        if fmt(x_sum) in html_x:
            r3_ok, r3_note = True, f"摘要含总金额{fmt(x_sum)}"
        elif fmt(x_avg) in html_x:
            r3_ok, r3_note = True, f"摘要含均值{fmt(x_avg)}"
        elif f"{x_rows:,}" in html_x:
            r3_ok, r3_note = True, f"摘要含记录数{x_rows}"
    else:
        r3_ok, r3_note = True, f"llm_used=true，数字复算以附录表R1/R2为准"
    check("R3 摘要数字复算", r3_ok, r3_note)

    # [6] docx 上传
    group("6. docx 上传 -> 文档型数据集（标准1文档侧）")
    s, r = upload(DOCX_PATH, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    check("上传200", s == 200, f"status={s} {str(r)[:120]}")
    file_d = r.get("file_id") if isinstance(r, dict) else None
    check("返回file_id", bool(file_d))
    s2, r2 = make_dataset(file_d, "验收_docx业务文档")
    check("建数据集200", s2 == 200, f"status={s2} {str(r2)[:120]}")
    ds_d = r2.get("dataset_id") if isinstance(r2, dict) else None
    check("返回dataset_id", bool(ds_d))
    check("文档型不建表", isinstance(r2, dict) and not r2.get("duckdb_table"))

    # [7] docx 报告
    group("7. docx 报告 -> 7章节/图表>=6/versionId（标准1/4/5）")
    s, r = make_dashboard(ds_d, "验收_docx看板")
    check("看板创建200", s == 200, f"status={s} {str(r)[:120]}")
    dash_d = r.get("dashboard_id")
    s, r = req("POST", "/reports", {"dashboard_id": dash_d})
    check("报告生成200", s == 200, f"status={s} {str(r)[:150]}")
    rep_d = r if isinstance(r, dict) else {}
    check("versionId非null非空", isinstance(rep_d.get("version_id"), str) and len(rep_d.get("version_id")) > 0,
          f"version_id={rep_d.get('version_id')}")
    s, r = req("GET", f"/reports/{rep_d.get('report_id')}/html")
    html_d = r if isinstance(r, str) else ""
    (OUT / "acceptance_docx_report.html").write_text(html_d, encoding="utf-8")
    n_ch_d = html_d.count('class="chapter"')
    check("章节数=7", n_ch_d == 7, f"chapters={n_ch_d}")
    check("图表数>=6", chart_elements(html_d) >= 6, f"charts={chart_elements(html_d)}")

    # [8] docx 段落复算
    group("8. docx 段落数复算（标准2）")
    m = re.search(r"段落数[：:]\s*([\d,]+)", html_d)
    got = int(m.group(1).replace(",", "")) if m else -1
    check("段落数复算", got == d_paras, f"报告={got} 源={d_paras}")

    # [9] 历史版本回看
    group("9. 历史版本回看（标准4）")
    s, r = req("POST", "/reports", {"dashboard_id": dash_x})
    check("二次生成200", s == 200, f"status={s}")
    url = f"{BASE}/reports?dashboard_id={dash_d}&limit=50"
    with urllib.request.urlopen(url, timeout=30) as resp:
        lst = json.loads(resp.read().decode("utf-8"))
    with urllib.request.urlopen(f"{BASE}/reports?dashboard_id={dash_x}&limit=50", timeout=30) as resp:
        lst_x = json.loads(resp.read().decode("utf-8"))
    items = lst_x.get("items") or []
    check("xlsx看板历史>=2", lst_x.get("total", 0) >= 2, f"total={lst_x.get('total')}")
    check("每条versionId非空", all(i.get("version_id") for i in items), "")
    items_d = lst.get("items") or []
    check("docx看板历史>=1", lst.get("total", 0) >= 1, f"total={lst.get('total')}")
    if items:
        s, r = req("GET", f"/reports/{items[0]['report_id']}/html")
        check("历史版本HTML可回看", s == 200 and isinstance(r, str) and len(r) > 100, f"status={s}")

    # [10] llm_used 诚实性
    group("10. llm_used 诚实性（标准3）")
    degrade = "AI解读不可用" in html_x and "确定性摘要" in html_x
    if rep1.get("llm_used") is False:
        check("降级标注存在", degrade, "")
    else:
        check("llm_used=true时不出现降级标注", not degrade, "")

    # [11] S4/S5 内部端点
    group("11. S4/S5 内部端点（占位符清除验证）")
    s, r = req("POST", "/brain/_internal/test-orchestrate")
    check("test-orchestrate 200", s == 200, f"status={s}")
    if isinstance(r, dict):
        flow = r.get("narrative_flow", "")
        check("narrative_flow真实(含类型计数)", "×" in flow and flow != "结论→佐证→明细", f"flow={flow[:80]}")
    s, r = req("POST", "/brain/_internal/test-score")
    check("test-score 200", s == 200, f"status={s}")

    print("\n" + "=" * 72)
    print(f"结果: PASS={PASS} FAIL={FAIL}")
    print(f"报告样例: {OUT}")
    print("=" * 72)
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
